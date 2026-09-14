#!/usr/bin/env python3
"""将定制简历Markdown批量转换为Word文档（优化排版版）。"""
import re
import sqlite3
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

PROJECT_ROOT = Path("/Volumes/固态硬盘1T/002-探索项目/040-OpenBiliClaw")
RESUME_DIR = PROJECT_ROOT / "求职知识库" / "03_岗位弹药库" / "定制简历"
OUTPUT_DIR = RESUME_DIR / "Word版"
OUTPUT_DIR.mkdir(exist_ok=True)


def add_run_with_bold(paragraph, text, font_size=10.5):
    """处理**加粗**标记。"""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(part)
        run.font.size = Pt(font_size)


def set_para_spacing(p, before=0, after=4, line_spacing=1.35):
    """统一设置段落间距。"""
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line_spacing
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def md_to_docx(md_path: Path, output_path: Path):
    doc = Document()

    # 设置默认字体和行间距
    style = doc.styles["Normal"]
    font = style.font
    font.name = "PingFang SC"
    font.size = Pt(10.5)
    pf = style.paragraph_format
    pf.line_spacing = 1.35
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.space_after = Pt(4)

    # 页边距
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    lines = md_path.read_text(encoding="utf-8").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line:
            i += 1
            continue

        # H1: 姓名
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line[2:])
            run.bold = True
            run.font.size = Pt(20)
            run.font.color.rgb = RGBColor(0x1A, 0x1B, 0x1C)
            set_para_spacing(p, before=0, after=8, line_spacing=1.2)
            i += 1
            continue

        # H2: 大章节
        if line.startswith("## "):
            p = doc.add_paragraph()
            run = p.add_run(line[3:])
            run.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0x2D, 0x5B, 0x9E)
            set_para_spacing(p, before=14, after=6, line_spacing=1.2)
            i += 1
            continue

        # H3: 工作经历/项目
        if line.startswith("### "):
            p = doc.add_paragraph()
            run = p.add_run(line[4:])
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            set_para_spacing(p, before=10, after=4, line_spacing=1.25)
            i += 1
            continue

        # 水平线
        if line.startswith("---"):
            p = doc.add_paragraph()
            set_para_spacing(p, before=4, after=4, line_spacing=1.0)
            i += 1
            continue

        # 无序列表
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_run_with_bold(p, line[2:], font_size=10.5)
            set_para_spacing(p, before=1, after=3, line_spacing=1.35)
            i += 1
            continue

        # 普通段落（联系方式、项目描述等）
        p = doc.add_paragraph()
        add_run_with_bold(p, line, font_size=10.5)
        set_para_spacing(p, before=2, after=4, line_spacing=1.35)
        i += 1

    doc.save(str(output_path))
    print(f"  ✓ {output_path.name}")


def main():
    conn = sqlite3.connect(str(PROJECT_ROOT / "data" / "resume.db"))
    c = conn.cursor()
    c.execute("SELECT id, company, target_position, version_name, file_path FROM resume_texts ORDER BY id")
    resumes = c.fetchall()
    conn.close()

    print(f"共 {len(resumes)} 份简历，开始转换（优化排版）...")
    for r in resumes:
        rid, company, position, version, file_path = r
        md_path = PROJECT_ROOT / file_path
        if not md_path.exists():
            print(f"  ✗ 文件不存在: {file_path}")
            continue
        safe_name = version.replace("/", "-").replace(" ", "")
        output_path = OUTPUT_DIR / f"童力-{safe_name}.docx"
        md_to_docx(md_path, output_path)

    print(f"\n完成！Word简历保存在: {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
