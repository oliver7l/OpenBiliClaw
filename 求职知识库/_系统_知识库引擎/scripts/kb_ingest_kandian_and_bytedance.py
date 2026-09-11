#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把两份新文件入库到面试知识库：
  1) 看点小说深度用户调研报告.docx -> 看点库.db (category=工作资料)
  2) 2026年09月03日-字节面试.zip  ->
       - 童力-百度-数据分析.pdf      -> 面试资料总库.db (category=岗位弹药)
       - 3 个 .url 飞书链接          -> 合成「字节面试-参考链接.md」一并入库 (category=岗位弹药)
入库前自动备份两个目标库到 data/_archive/。
依赖：python-docx, pypdf, sqlite_vec, kb_migrate_p1.chunk_text
"""
import os, sys, re, json, shutil, sqlite3, sqlite_vec, zipfile, urllib.request
from datetime import datetime

ROOT = "/Volumes/固态硬盘1T/002-探索项目/040-OpenBiliClaw"
SCRIPTS = os.path.join(ROOT, "求职知识库/_系统_知识库引擎/scripts")
DATA = os.path.join(ROOT, "求职知识库/_系统_知识库引擎/数据")
SRC_ROOT = os.path.join(ROOT, "求职知识库/01_原始资料库")
sys.path.insert(0, SCRIPTS)
from kb_migrate_p1 import chunk_text  # 标题感知分块

KANDIAN_DB = os.path.join(DATA, "看点库.db")
MASTER_DB = os.path.join(DATA, "面试资料总库.db")
MIN_CHARS = 60  # 小于此字数的块不向量化

EMBED_MODEL = "bge-m3"

def connect(db_path):
    conn = sqlite3.connect(db_path)
    conn.enable_load_extension(True)
    sqlite_vec.load(conn)
    return conn

def backup(conn, db_path, archive_dir):
    ts = datetime.now().strftime("%Y%m%d-%H%M%S")
    os.makedirs(archive_dir, exist_ok=True)
    dst = os.path.join(archive_dir, f"{os.path.basename(db_path)}.bak-{ts}")
    shutil.copy2(db_path, dst)
    print(f"[backup] {dst}")
    return dst

def embed(texts):
    """批量调用 Ollama bge-m3，返回 list[list[float]]，与输入等长。"""
    if not texts:
        return []
    payload = json.dumps({"model": EMBED_MODEL, "input": texts}).encode()
    req = urllib.request.Request("http://localhost:11434/api/embed",
                                 data=payload, headers={"Content-Type": "application/json"})
    resp = urllib.request.urlopen(req, timeout=120).read()
    data = json.loads(resp)
    return data["embeddings"]

def sanitize(t):
    t = t.replace("\u0000", "")
    t = re.sub(r"[\ud800-\udfff]", "", t)
    return t

def extract_docx(path):
    import docx
    d = docx.Document(path)
    parts = []
    for p in d.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    # 表格也抽出来，避免漏内容
    for tb in d.tables:
        for row in tb.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)

def extract_pdf(path):
    from pypdf import PdfReader
    r = PdfReader(path)
    return "\n".join((p.extract_text() or "") for p in r.pages)

def ingest_doc(conn, src_path, rel_path, category, title, text, archive_dir, phys_root):
    """抽取(已给 text) -> 拷贝归档 -> 写 doc/doc_content/doc_chunk/doc_vector。
    返回 (doc_id, n_chunks, n_vec)。doc_fts 由触发器自动维护。
    phys_root: 实体文件落盘根目录；rel_path 已含其下的相对路径（注意两库 rel_path 前缀不同）。"""
    # 1) 归档实体文件
    dest = os.path.join(phys_root, rel_path)
    os.makedirs(os.path.dirname(dest), exist_ok=True)
    if os.path.abspath(src_path) != os.path.abspath(dest):
        shutil.copy2(src_path, dest)
    text = sanitize(text)
    char_count = len(text)
    # 2) doc（触发器自动维护 doc_fts）
    cur = conn.execute(
        "INSERT INTO doc(path,rel_path,category,ext,size,mtime,sha,title,content,char_count,extract_method,status,updated_at) "
        "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (rel_path, rel_path, category,
         os.path.splitext(src_path)[1].lstrip(".").lower(),
         os.path.getsize(src_path), os.path.getmtime(src_path), "",
         title, text, char_count, "auto", "ok",
         datetime.now().isoformat(timespec="seconds")))
    doc_id = cur.lastrowid
    # 3) doc_content（全文）
    conn.execute(
        "INSERT INTO doc_content(doc_id,version,content,char_count,extract_method,created_at) "
        "VALUES(?,?,?,?,?,?)",
        (doc_id, 1, text, char_count, "auto", datetime.now().isoformat(timespec="seconds")))
    # 4) doc_chunk
    chunks = chunk_text(text)
    n_chunks = 0
    for seq, (header, t) in enumerate(chunks):
        conn.execute(
            "INSERT INTO doc_chunk(doc_id,content_version,seq,header,text,char_count) VALUES(?,?,?,?,?,?)",
            (doc_id, 1, seq, header, t, len(t)))
        n_chunks += 1
    # 5) doc_vector（仅 >= MIN_CHARS 的块；INSERT OR IGNORE 幂等）
    rows = conn.execute(
        "SELECT id,text FROM doc_chunk WHERE doc_id=? AND char_count>=?", (doc_id, MIN_CHARS)).fetchall()
    if rows:
        embs = embed([r[1] for r in rows])
        for (cid, _), emb in zip(rows, embs):
            conn.execute("INSERT OR IGNORE INTO doc_vector(chunk_id,embedding) VALUES(?,?)",
                         (cid, sqlite_vec.serialize_float32(emb)))
    conn.commit()
    return doc_id, n_chunks, len(rows)

def main():
    ts = datetime.now().strftime("%Y%m%d-%H%M%S")
    archive_dir = os.path.join(DATA, "_archive")

    # ---- 文件定位 ----
    DOCX = "/Users/imac/Downloads/看点小说深度用户调研报告.docx"
    ZIP = "/Users/imac/Downloads/2026年09月03日-字节面试.zip"
    assert os.path.exists(DOCX), DOCX
    assert os.path.exists(ZIP), ZIP

    # 自动备份两个目标库
    kc = connect(KANDIAN_DB); backup(kc, KANDIAN_DB, archive_dir); kc.close()
    mc = connect(MASTER_DB); backup(mc, MASTER_DB, archive_dir); mc.close()

    # ===== 1) docx -> 看点库.db =====
    print("\n========== [1/2] docx -> 看点库.db ==========")
    kc = connect(KANDIAN_DB)
    dtext = extract_docx(DOCX)
    docx_rel = "01_原始资料库/03_工作资料/工作资料_腾讯/2020年06月-看点小说深度调研/看点小说深度用户调研报告.docx"
    did, nch, nvec = ingest_doc(
        kc, DOCX, docx_rel, "工作资料",
        "看点小说深度用户调研报告", dtext, SRC_ROOT,
        os.path.join(ROOT, "求职知识库"))  # 看点库 rel_path 已含 01_原始资料库/ 前缀
    kc.close()
    print(f"  docx -> 看点库 doc_id={did} chunks={nch} vec={nvec} chars={len(dtext)}")

    # ===== 2) zip -> 面试资料总库.db =====
    print("\n========== [2/2] zip -> 面试资料总库.db (岗位弹药) ==========")
    tmp = f"/tmp/kb_zip_bytedance_{ts}"
    if os.path.exists(tmp):
        shutil.rmtree(tmp)
    os.makedirs(tmp)
    with zipfile.ZipFile(ZIP) as z:
        z.extractall(tmp)
    # 找到解压后的根目录
    inner = os.path.join(tmp, "2026年09月03日-字节面试")
    assert os.path.isdir(inner), f"zip 内目录缺失: {inner}"

    mc = connect(MASTER_DB)
    # 2a) PDF
    pdf_src = os.path.join(inner, "童力-百度-数据分析.pdf")
    ptext = extract_pdf(pdf_src)
    pdf_rel = "03_岗位弹药库/2026年09月03日-字节面试/童力-百度-数据分析.pdf"
    did1, nch1, nvec1 = ingest_doc(
        mc, pdf_src, pdf_rel, "岗位弹药",
        "童力-百度-数据分析（简历参考）", ptext, SRC_ROOT, SRC_ROOT)
    print(f"  pdf  -> 主库 doc_id={did1} chunks={nch1} vec={nvec1} chars={len(ptext)}")

    # 2b) 3 个 .url -> 合成一份 markdown 入库
    url_map = {}
    for fn in os.listdir(inner):
        if fn.endswith(".url"):
            p = os.path.join(inner, fn)
            raw = open(p, encoding="utf-8", errors="ignore").read()
            m = re.search(r"URL=(.+)", raw)
            if m:
                url = m.group(1).strip()
                label = fn[:-4]
                url_map[label] = url
    md_lines = ["# 字节面试 - 参考链接", "",
                "> 来源：2026年09月03日-字节面试.zip 内的飞书快捷方式（需登录飞书后查看）", ""]
    for label, url in url_map.items():
        md_lines.append(f"- **{label}**：{url}")
    md_text = "\n".join(md_lines)
    # 实体也归档：把 pdf + .url + 这份 md 都落进 岗位弹药目录
    md_rel = "03_岗位弹药库/2026年09月03日-字节面试/字节面试-参考链接.md"
    md_src = os.path.join(inner, "字节面试-参考链接.md")
    with open(md_src, "w", encoding="utf-8") as f:
        f.write(md_text)
    did2, nch2, nvec2 = ingest_doc(
        mc, md_src, md_rel, "岗位弹药",
        "字节面试-参考链接（飞书）", md_text, SRC_ROOT, SRC_ROOT)
    print(f"  links-> 主库 doc_id={did2} chunks={nch2} vec={nvec2} urls={len(url_map)}")
    # 同时把原始 .url 文件也归档进该目录
    for fn in os.listdir(inner):
        if fn.endswith(".url"):
            shutil.copy2(os.path.join(inner, fn),
                         os.path.join(SRC_ROOT, os.path.dirname(md_rel), fn))
    mc.close()

    print("\n=== DONE ===")
    print(f"看点库: docx 已入 (doc_id={did})")
    print(f"主库  : PDF(doc_id={did1}) + 链接MD(doc_id={did2}) 已入，均 category=岗位弹药")

if __name__ == "__main__":
    main()
